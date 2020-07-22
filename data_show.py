import json
from openpyxl import *
from openpyxl.chart import (LineChart, Reference)
from os import (listdir, mkdir, path, remove)
import pandas as pd
import requests




def get_markets():
    
    data = bytes_to_str(read_xml('https://www.predictit.org/api/marketdata/all'))
   
    data = json.loads(data)
    markets = data['markets']  
    
    return markets


def read_xml(url):

    r = requests.get(url)

    return r.content


def bytes_to_str(data):

    return data.decode()


def to_xl():
    if (not path.exists("markets")):
        mkdir("markets")

    markets = get_markets()

    for market in markets:
        market['shortName'] = market['shortName'].replace("?", "_").replace("/", "-").replace("'", "").replace("\"", "")
        filepath = 'markets/' + market['shortName'] + '.xlsx'
        if (path.exists(filepath)):
            
            xl = pd.read_excel(filepath)
        
            contracts = market['contracts']
            df = pd.DataFrame(contracts)
            df['timeStamp'] = market['timeStamp']
            del df['id']
            del df['name']
            del df['status']
            del df['dateEnd']
            del df['image']
            del df['displayOrder']
            
            df_new = pd.concat([df, xl])
            df_new = df_new.sort_values(by=['shortName', 'timeStamp'])
            
            remove(filepath)
           
            writer = pd.ExcelWriter(filepath, engine='openpyxl')
            df_new.to_excel(writer, index=False)
            
            writer.save()
            writer.close()

        else:
            
            writer = pd.ExcelWriter(filepath, engine='openpyxl')

            contracts = market['contracts']
            df = pd.DataFrame(contracts)
            df['timeStamp'] = market['timeStamp']
            del df['id']
            del df['name']
            del df['status']
            del df['dateEnd']
            del df['image']
            del df['displayOrder']
            df.to_excel(writer, index=False)
           
            writer.save()
            writer.close()
    

def create_charts():
    to_xl()
    excels = listdir('markets')
    
    for excel in excels:
        filepath = 'markets/' + excel
        writer = pd.ExcelWriter(filepath , engine='openpyxl')
        df = pd.read_excel(filepath)
        workbook = load_workbook(filepath)
        writer.book = workbook
        
        worksheet=writer.book.worksheets[0]
        total_elems = 0
        
        dups = df.pivot_table(index=['shortName'], aggfunc='size')
        unique_names = df.shortName.unique()
        
        if (len(dups) == 1):
            chart = LineChart()
            chart.title = df['shortName'][0]
            chart.y_axis.title = 'price'
            chart.x_axis.title = 'time stamp'
            
            data = Reference(worksheet, min_col=2, min_row=2, max_col=7, max_row=dups)
            
            chart.add_data(data)
            
            worksheet.add_chart(chart, 'L4')
        else:
            i = 2
            j = 0
            for dup in dups:
                total_elems += dup
                
                chart = LineChart()
                chart.title = unique_names[j]
                j += 1
                
                chart.y_axis.title = 'price'
                chart.x_axis.title = 'time stamp'
                
                data = Reference(worksheet, min_col=2, min_row=i, max_col=7, max_row=total_elems+1)
                chart.add_data(data)
                
                worksheet.add_chart(chart, 'L'+str(i*4))
                i += dup
                
        writer.save()
        writer.close()


create_charts()